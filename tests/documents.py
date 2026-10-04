"""Build small but real PDF, PowerPoint and Word files for tests."""

from __future__ import annotations

from pathlib import Path


def make_pdf(path: Path, pages: list[list[str]]) -> Path:
    """A text PDF: one list of lines per page (Helvetica, no dependencies)."""
    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    objects: list[bytes] = []

    def add(body: str | bytes) -> int:
        objects.append(body.encode("latin-1") if isinstance(body, str) else body)
        return len(objects)

    catalog = add("")  # filled in below
    pages_id = add("")
    font = add("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    kids = []
    for lines in pages:
        commands = ["BT", "/F1 12 Tf", "72 770 Td"]
        for n, line in enumerate(lines):
            if n:
                commands.append("0 -16 Td")
            commands.append(f"({escape(line)}) Tj")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        content = add(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
        kids.append(add(f"<< /Type /Page /Parent {pages_id} 0 R /MediaBox [0 0 595 842] "
                        f"/Resources << /Font << /F1 {font} 0 R >> >> /Contents {content} 0 R >>"))
    objects[catalog - 1] = f"<< /Type /Catalog /Pages {pages_id} 0 R >>".encode()
    objects[pages_id - 1] = (f"<< /Type /Pages /Kids [{' '.join(f'{k} 0 R' for k in kids)}] "
                             f"/Count {len(kids)} >>").encode()

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root {catalog} 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    path.write_bytes(bytes(out))
    return path


def make_pptx(path: Path, slides: list[tuple[str, str]], notes: dict[int, str] | None = None) -> Path:
    """A deck of (title, body) slides; ``notes`` maps slide index to speaker notes."""
    from pptx import Presentation

    deck = Presentation()
    for index, (title, body) in enumerate(slides):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
        if notes and index in notes:
            slide.notes_slide.notes_text_frame.text = notes[index]
    deck.save(str(path))
    return path


def make_docx(path: Path, sections: list[tuple[str, list[str]]]) -> Path:
    """A Word file with a Heading 1 and paragraphs per section ('' title = no heading)."""
    import docx

    document = docx.Document()
    for title, paragraphs in sections:
        if title:
            document.add_heading(title, level=1)
        for paragraph in paragraphs:
            document.add_paragraph(paragraph)
    document.save(str(path))
    return path
